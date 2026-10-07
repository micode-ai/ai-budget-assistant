"""Write the downloadable household-budget workbook offered by topic 28 (pair `excel-budget`).

One file per language, `assets/<lang>/budget-template.xlsx`, laid out exactly as the article
describes: a Transactions sheet with five columns, a fixed category list, and a Summary sheet
of SUMIFS by category x month with a running balance. build_blog.py copies `assets/` into the
built site at /blog/<lang>/assets/.

Run by hand when the layout or a translation changes — NOT from build_blog.py: xlsx is a zip
with timestamps, so regenerating it on every build would commit a changed binary each time.

    python docs/marketing/seo/build_excel_template.py
"""
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.abspath(__file__))
FIXED = datetime(2026, 10, 7)

T = {
    "en": {"tx": "Transactions", "sum": "Summary", "cats": "Categories",
           "cols": ["Date", "Category", "Description", "Amount", "Account"],
           "income": "Income", "total": "Total spent", "balance": "Balance (income - spent)",
           "start": "Starting balance", "category": "Category",
           "months": ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"],
           "list": ["Housing", "Groceries", "Eating out", "Transport", "Utilities", "Health",
                    "Subscriptions", "Shopping", "Kids", "Savings", "Debt payments", "Fun"],
           "note": "Log expenses as positive amounts. Use the category 'Income' for money coming in.",
           "app": "Tired of typing every row? AI Budget Assistant logs expenses by voice or receipt photo: https://ai-budget.pl/"},
    "pl": {"tx": "Transakcje", "sum": "Podsumowanie", "cats": "Kategorie",
           "cols": ["Data", "Kategoria", "Opis", "Kwota", "Konto"],
           "income": "Dochód", "total": "Wydatki razem", "balance": "Saldo (dochód - wydatki)",
           "start": "Saldo początkowe", "category": "Kategoria",
           "months": ["Sty", "Lut", "Mar", "Kwi", "Maj", "Cze", "Lip", "Sie", "Wrz", "Paź", "Lis", "Gru"],
           "list": ["Mieszkanie", "Jedzenie", "Jedzenie na mieście", "Transport", "Rachunki", "Zdrowie",
                    "Subskrypcje", "Zakupy", "Dzieci", "Oszczędności", "Spłata długów", "Rozrywka"],
           "note": "Wydatki wpisuj jako kwoty dodatnie. Wpływy zapisuj z kategorią „Dochód”.",
           "app": "Masz dość wpisywania każdego wiersza? AI Budget Assistant zapisuje wydatki głosem lub ze zdjęcia paragonu: https://ai-budget.pl/"},
    "de": {"tx": "Buchungen", "sum": "Übersicht", "cats": "Kategorien",
           "cols": ["Datum", "Kategorie", "Beschreibung", "Betrag", "Konto"],
           "income": "Einnahmen", "total": "Ausgaben gesamt", "balance": "Saldo (Einnahmen - Ausgaben)",
           "start": "Anfangssaldo", "category": "Kategorie",
           "months": ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"],
           "list": ["Wohnen", "Lebensmittel", "Auswärts essen", "Mobilität", "Nebenkosten", "Gesundheit",
                    "Abos", "Einkäufe", "Kinder", "Sparen", "Schuldentilgung", "Freizeit"],
           "note": "Ausgaben als positive Beträge eintragen. Geldeingänge mit der Kategorie „Einnahmen“ erfassen.",
           "app": "Keine Lust mehr, jede Zeile zu tippen? AI Budget Assistant erfasst Ausgaben per Sprache oder Kassenbon-Foto: https://ai-budget.pl/de/"},
    "es": {"tx": "Movimientos", "sum": "Resumen", "cats": "Categorías",
           "cols": ["Fecha", "Categoría", "Descripción", "Importe", "Cuenta"],
           "income": "Ingresos", "total": "Total gastado", "balance": "Saldo (ingresos - gastos)",
           "start": "Saldo inicial", "category": "Categoría",
           "months": ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"],
           "list": ["Vivienda", "Supermercado", "Comer fuera", "Transporte", "Suministros", "Salud",
                    "Suscripciones", "Compras", "Hijos", "Ahorro", "Pago de deudas", "Ocio"],
           "note": "Anota los gastos como importes positivos. Usa la categoría «Ingresos» para el dinero que entra.",
           "app": "¿Cansado de escribir cada fila? AI Budget Assistant registra gastos por voz o con una foto del ticket: https://ai-budget.pl/es/"},
    "fr": {"tx": "Transactions", "sum": "Synthèse", "cats": "Catégories",
           "cols": ["Date", "Catégorie", "Description", "Montant", "Compte"],
           "income": "Revenus", "total": "Total dépensé", "balance": "Solde (revenus - dépenses)",
           "start": "Solde initial", "category": "Catégorie",
           "months": ["Janv", "Févr", "Mars", "Avr", "Mai", "Juin", "Juil", "Août", "Sept", "Oct", "Nov", "Déc"],
           "list": ["Logement", "Courses", "Restaurants", "Transport", "Charges", "Santé",
                    "Abonnements", "Achats", "Enfants", "Épargne", "Remboursement de dettes", "Loisirs"],
           "note": "Saisissez les dépenses en montants positifs. Utilisez la catégorie « Revenus » pour l'argent qui entre.",
           "app": "Marre de saisir chaque ligne ? AI Budget Assistant enregistre vos dépenses à la voix ou en photo de ticket : https://ai-budget.pl/fr/"},
    "nl": {"tx": "Transacties", "sum": "Overzicht", "cats": "Categorieën",
           "cols": ["Datum", "Categorie", "Omschrijving", "Bedrag", "Rekening"],
           "income": "Inkomsten", "total": "Totaal uitgegeven", "balance": "Saldo (inkomsten - uitgaven)",
           "start": "Beginsaldo", "category": "Categorie",
           "months": ["Jan", "Feb", "Mrt", "Apr", "Mei", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dec"],
           "list": ["Wonen", "Boodschappen", "Uit eten", "Vervoer", "Vaste lasten", "Gezondheid",
                    "Abonnementen", "Winkelen", "Kinderen", "Sparen", "Schulden aflossen", "Vrije tijd"],
           "note": "Vul uitgaven in als positieve bedragen. Gebruik de categorie 'Inkomsten' voor geld dat binnenkomt.",
           "app": "Geen zin meer om elke regel te typen? AI Budget Assistant legt uitgaven vast met je stem of een foto van de bon: https://ai-budget.pl/nl/"},
    "ru": {"tx": "Операции", "sum": "Сводка", "cats": "Категории",
           "cols": ["Дата", "Категория", "Описание", "Сумма", "Счёт"],
           "income": "Доход", "total": "Всего потрачено", "balance": "Остаток (доход - расходы)",
           "start": "Начальный остаток", "category": "Категория",
           "months": ["Янв", "Фев", "Мар", "Апр", "Май", "Июн", "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек"],
           "list": ["Жильё", "Продукты", "Кафе и рестораны", "Транспорт", "Коммунальные", "Здоровье",
                    "Подписки", "Покупки", "Дети", "Накопления", "Погашение долгов", "Развлечения"],
           "note": "Расходы записывайте положительными суммами. Поступления — с категорией «Доход».",
           "app": "Надоело вбивать каждую строку? AI Budget Assistant записывает расходы голосом или по фото чека: https://ai-budget.pl/ru/"},
    "ua": {"tx": "Операції", "sum": "Підсумок", "cats": "Категорії",
           "cols": ["Дата", "Категорія", "Опис", "Сума", "Рахунок"],
           "income": "Дохід", "total": "Усього витрачено", "balance": "Залишок (дохід - витрати)",
           "start": "Початковий залишок", "category": "Категорія",
           "months": ["Січ", "Лют", "Бер", "Кві", "Тра", "Чер", "Лип", "Сер", "Вер", "Жов", "Лис", "Гру"],
           "list": ["Житло", "Продукти", "Кафе й ресторани", "Транспорт", "Комунальні", "Здоров'я",
                    "Підписки", "Покупки", "Діти", "Заощадження", "Погашення боргів", "Розваги"],
           "note": "Витрати записуйте додатними сумами. Надходження — з категорією «Дохід».",
           "app": "Набридло вводити кожен рядок? AI Budget Assistant записує витрати голосом або за фото чека: https://ai-budget.pl/ua/"},
    "be": {"tx": "Аперацыі", "sum": "Зводка", "cats": "Катэгорыі",
           "cols": ["Дата", "Катэгорыя", "Апісанне", "Сума", "Рахунак"],
           "income": "Даход", "total": "Усяго выдаткавана", "balance": "Рэшта (даход - выдаткі)",
           "start": "Пачатковая рэшта", "category": "Катэгорыя",
           "months": ["Сту", "Лют", "Сак", "Кра", "Тра", "Чэр", "Ліп", "Жні", "Вер", "Кас", "Ліс", "Сне"],
           "list": ["Жыллё", "Прадукты", "Кавярні і рэстараны", "Транспарт", "Камунальныя", "Здароўе",
                    "Падпіскі", "Пакупкі", "Дзеці", "Зберажэнні", "Пагашэнне даўгоў", "Забавы"],
           "note": "Выдаткі запісвайце дадатнымі сумамі. Паступленні — з катэгорыяй «Даход».",
           "app": "Надакучыла ўводзіць кожны радок? AI Budget Assistant запісвае выдаткі голасам або па фота чэка: https://ai-budget.pl/be/"},
}

HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="F5832A")
ROWS = 1000  # pre-formatted transaction rows the Summary formulas cover


def build(lang):
    t = T[lang]
    wb = Workbook()
    tx = wb.active
    tx.title = t["tx"]
    tx.append(t["cols"])
    for c in range(1, 6):
        cell = tx.cell(row=1, column=c)
        cell.font, cell.fill = HEAD, FILL
    for col, w in zip("ABCDE", (12, 22, 36, 12, 16)):
        tx.column_dimensions[col].width = w
    tx.freeze_panes = "A2"
    for r in range(2, ROWS + 2):
        tx.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        tx.cell(row=r, column=4).number_format = "#,##0.00"
    tx.cell(row=1, column=7, value=t["note"]).font = Font(italic=True, color="777777")

    cats = wb.create_sheet(t["cats"])
    cats.append([t["category"]])
    cats["A1"].font, cats["A1"].fill = HEAD, FILL
    for name in t["list"] + [t["income"]]:
        cats.append([name])
    cats.column_dimensions["A"].width = 26
    n = len(t["list"]) + 1
    dv = DataValidation(type="list", formula1=f"='{t['cats']}'!$A$2:$A${n + 1}", allow_blank=True)
    tx.add_data_validation(dv)
    dv.add(f"B2:B{ROWS + 1}")

    sm = wb.create_sheet(t["sum"], 0)
    year = 2026
    sm.cell(row=1, column=1, value=t["category"])
    for m, label in enumerate(t["months"], start=2):
        sm.cell(row=1, column=m, value=f"{label} {year}")
    for c in range(1, 14):
        cell = sm.cell(row=1, column=c)
        cell.font, cell.fill = HEAD, FILL
        cell.alignment = Alignment(horizontal="center")
    sm.column_dimensions["A"].width = 28
    rng = lambda col: f"'{t['tx']}'!${col}$2:${col}${ROWS + 1}"
    for i, name in enumerate(t["list"] + [t["income"]], start=2):
        sm.cell(row=i, column=1, value=name)
        for m in range(1, 13):
            col = get_column_letter(m + 1)
            start = f"DATE({year},{m},1)"
            end = f"EDATE({start},1)"
            sm[f"{col}{i}"] = (f'=SUMIFS({rng("D")},{rng("B")},$A{i},'
                               f'{rng("A")},">="&{start},{rng("A")},"<"&{end})')
            sm[f"{col}{i}"].number_format = "#,##0.00"
    last_cat = len(t["list"]) + 1
    inc_row = last_cat + 1
    tot, bal, startr = inc_row + 2, inc_row + 3, inc_row + 5
    sm.cell(row=tot, column=1, value=t["total"]).font = Font(bold=True)
    sm.cell(row=bal, column=1, value=t["balance"]).font = Font(bold=True)
    sm.cell(row=startr, column=1, value=t["start"])
    sm.cell(row=startr, column=2, value=0).number_format = "#,##0.00"
    for m in range(1, 13):
        col = get_column_letter(m + 1)
        prev = get_column_letter(m)
        sm[f"{col}{tot}"] = f"=SUM({col}2:{col}{last_cat})"
        carried = f"$B${startr}" if m == 1 else f"{prev}{bal}"
        sm[f"{col}{bal}"] = f"={carried}+{col}{inc_row}-{col}{tot}"
        for r in (tot, bal):
            sm[f"{col}{r}"].number_format = "#,##0.00"
            sm[f"{col}{r}"].font = Font(bold=True)
    sm.cell(row=startr + 2, column=1, value=t["app"]).font = Font(italic=True, color="777777")
    sm.freeze_panes = "B2"

    wb.properties.creator = "AI Budget Assistant"
    wb.properties.created = wb.properties.modified = FIXED
    out = os.path.join(ROOT, "assets", lang)
    os.makedirs(out, exist_ok=True)
    wb.save(os.path.join(out, "budget-template.xlsx"))


if __name__ == "__main__":
    for lang in T:
        build(lang)
        print("wrote", lang)
